"""Export Viva Defense Guide and Week 3 Progress Summary to Word docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def _setup_document_styles(doc: Document) -> None:
    """Configure document margins and default typography."""
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    normal_style = doc.styles["Normal"]
    font = normal_style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.25


def _set_cell_shading(cell: _Cell, color_hex: str) -> None:
    """Apply background shading color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tc_pr.append(shd)


def _set_cell_left_border(
    cell: _Cell, color_hex: str = "1864AB", sz: str = "24"
) -> None:
    """Apply a thick left border to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f"</w:tcBorders>"
    )
    tc_pr.append(borders)


def _set_box_borders(
    cell: _Cell, color_hex: str = "CED4DA", sz: str = "4"
) -> None:
    """Set thin borders around all sides of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f"</w:tcBorders>"
    )
    tc_pr.append(borders)


def _set_table_borders(table: Table, color_hex: str = "CED4DA") -> None:
    """Set light gray grid borders for a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f"</w:tblBorders>"
    )
    tbl_pr.append(borders)


def _apply_token_formatting(
    run: Run,
    token: str,
    default_bold: bool = False,
    default_italic: bool = False,
) -> None:
    """Apply inline bold, italic, or monospace styling to a run."""
    run.bold = default_bold
    run.italic = default_italic
    if token.startswith("***") and token.endswith("***") and len(token) >= 6:
        run.text = token[3:-3]
        run.bold = True
        run.italic = True
    elif token.startswith("**") and token.endswith("**") and len(token) >= 4:
        run.text = token[2:-2]
        run.bold = True
    elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
        run.text = token[1:-1]
        run.italic = True
    elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
        run.text = token[1:-1]
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
    else:
        run.text = token


def _add_inline_formatted_text(
    paragraph: Paragraph,
    text: str,
    default_bold: bool = False,
    default_italic: bool = False,
) -> None:
    """Parse inline markdown tokens and add formatted runs to a paragraph."""
    pattern = re.compile(r"(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`[^`]+`)")
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        _apply_token_formatting(
            run, token, default_bold=default_bold, default_italic=default_italic
        )


def _add_heading(
    doc: Document, text: str, level: int, is_first_h1: bool = False
) -> None:
    """Add a styled heading paragraph to the document."""
    style_level = min(level, 3)
    p = doc.add_paragraph(style=f"Heading {style_level}")
    if level == 1 and not is_first_h1:
        p.paragraph_format.page_break_before = True

    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True

    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(11, 59, 96)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(24, 100, 171)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(43, 58, 66)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)


def _add_blockquote(doc: Document, lines: list[str]) -> None:
    """Add a blockquote callout box with gray background and left border."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F1F3F5")
    _set_cell_left_border(cell, color_hex="1864AB", sz="24")

    first = True
    for raw in lines:
        cleaned = re.sub(r"^>\s?", "", raw).strip()
        if not cleaned:
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        _add_inline_formatted_text(p, cleaned, default_italic=True)


def _add_code_block(doc: Document, code_lines: list[str]) -> None:
    """Add a monospace code block in a bordered gray box."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F8F9FA")
    _set_box_borders(cell, color_hex="CED4DA", sz="4")

    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(33, 37, 41)


def _parse_table_row(line: str) -> list[str]:
    """Parse cell values from a markdown table row."""
    parts = [c.strip() for c in line.split("|")]
    if parts and parts[0] == "":
        parts = parts[1:]
    if parts and parts[-1] == "":
        parts = parts[:-1]
    return parts


def _add_table(doc: Document, table_lines: list[str]) -> None:
    """Convert markdown table lines into a styled Word table."""
    rows_data: list[list[str]] = []
    for line in table_lines:
        if re.match(r"^\|?[\s:-|]+\|?$", line.strip()) and "-" in line:
            continue
        cells = _parse_table_row(line)
        if cells:
            rows_data.append(cells)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    _set_table_borders(table, color_hex="CED4DA")

    for row_idx, row_content in enumerate(rows_data):
        is_header = row_idx == 0
        for col_idx in range(num_cols):
            cell_text = row_content[col_idx] if col_idx < len(row_content) else ""
            cell = table.cell(row_idx, col_idx)
            if is_header:
                _set_cell_shading(cell, "E7F5FF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            _add_inline_formatted_text(p, cell_text, default_bold=is_header)


def _consume_code_block(
    lines: list[str], start_idx: int
) -> tuple[list[str], int]:
    """Collect lines inside a fenced markdown code block."""
    collected: list[str] = []
    idx = start_idx + 1
    while idx < len(lines):
        if lines[idx].strip().startswith("```"):
            return collected, idx + 1
        collected.append(lines[idx])
        idx += 1
    return collected, idx


def _consume_prefixed_block(
    lines: list[str], start_idx: int, prefix: str
) -> tuple[list[str], int]:
    """Collect contiguous lines that begin with a specific prefix."""
    collected: list[str] = []
    idx = start_idx
    while idx < len(lines) and lines[idx].strip().startswith(prefix):
        collected.append(lines[idx])
        idx += 1
    return collected, idx


def _handle_heading_line(
    doc: Document, stripped: str, first_h1: bool, idx: int
) -> tuple[bool, int]:
    """Parse heading prefix and text, then add heading paragraph."""
    level = len(stripped) - len(stripped.lstrip("#"))
    heading_text = stripped.lstrip("#").strip()
    is_first = level == 1 and first_h1
    _add_heading(doc, heading_text, level, is_first_h1=is_first)
    return (False if level == 1 else first_h1), idx + 1


def _add_body_paragraph(doc: Document, text: str) -> None:
    """Add a bullet point or standard body paragraph."""
    if text.startswith(("- ", "* ")):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2)
        _add_inline_formatted_text(p, text[2:].strip())
        return

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    _add_inline_formatted_text(p, text)


def _process_markdown_lines(doc: Document, lines: list[str]) -> None:
    """Iterate through markdown lines and dispatch each block to Word."""
    idx = 0
    first_h1 = True
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        if not stripped or stripped == "---":
            idx += 1
            continue
        if stripped.startswith("```"):
            code_lines, idx = _consume_code_block(lines, idx)
            _add_code_block(doc, code_lines)
        elif stripped.startswith(">"):
            bq_lines, idx = _consume_prefixed_block(lines, idx, ">")
            _add_blockquote(doc, bq_lines)
        elif stripped.startswith("|"):
            tbl_lines, idx = _consume_prefixed_block(lines, idx, "|")
            _add_table(doc, tbl_lines)
        elif stripped.startswith("#"):
            first_h1, idx = _handle_heading_line(doc, stripped, first_h1, idx)
        else:
            _add_body_paragraph(doc, stripped)
            idx += 1


def _resolve_input_path(input_md: str | Path) -> Path:
    """Resolve input path relative to current dir or repository root."""
    path = Path(input_md)
    if path.exists():
        return path
    repo_fallback = Path(__file__).resolve().parent.parent / input_md
    if repo_fallback.exists():
        return repo_fallback
    return path


def convert_viva_guide_to_docx(
    input_md: str | Path, output_docx: str | Path
) -> Path:
    """Convert viva defense markdown guide to a formatted Word document."""
    in_path = _resolve_input_path(input_md)
    out_path = Path(output_docx)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    content = in_path.read_text(encoding="utf-8")
    doc = Document()
    _setup_document_styles(doc)
    _process_markdown_lines(doc, content.splitlines())
    doc.save(str(out_path))
    return out_path


def _build_week3_summary_header(doc: Document) -> None:
    """Add title and academic metadata for Week 3 report."""
    _add_heading(
        doc,
        "BÁO CÁO TIẾN ĐỘ TUẦN 3: ĐỀ TÀI NIÊN LUẬN KỸ THUẬT PHẦN MỀM (CT250H)",
        level=1,
        is_first_h1=True,
    )
    _add_heading(
        doc,
        "Temporal Knowledge Graph-Augmented Generation for Vietnamese VAT Regulations",
        level=2,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _add_inline_formatted_text(
        p,
        "**Đơn vị:** Khoa CNTT&TT - Trường Đại học Cần Thơ | **Mốc:** Tuần 3/8 | "
        "**Thành viên:** Nhóm K49 Kỹ thuật Phần mềm CLC",
    )


def _build_week3_goals_section(doc: Document) -> None:
    """Add Week 3 sprint goals and engineering deliverables table."""
    _add_heading(doc, "1. Mục Tiêu & Kết Quả Nghiệm Thu Sprint 1–3", level=2)
    tbl_lines = [
        "| Trụ cột Kỹ thuật | Kết quả Nghiệm thu | Trạng thái |",
        "| :--- | :--- | :--- |",
        "| **Dữ liệu Pháp lý sạch** | 5 văn bản QPPL, 76 Điều, 271 Khoản, 325 Điểm qua Parser 2 cấp | Hoàn thành 100% |",
        "| **Đồ thị Neo4j Đa quan hệ** | 2 tầng: Cây phân cấp + 8 AMENDS, 3 GUIDES, 17 CITES | Hoàn thành 100% |",
        "| **Tìm kiếm lai SOTA** | Fulltext Lucene + Dense Vector BGE-M3 (1024 chiều) + RRF (k=60) | Hoàn thành 100% |",
        "| **Tiền lọc Thời gian tất định** | RANGE INDEX valid_from/to với Sentinel Date '9999-12-31' | Cypher < 5ms |",
        "| **Prototype Kiểm thử** | Streamlit Split-View kết nối trực tiếp Neo4j 5.26 | Hoàn thành 100% |",
    ]
    _add_table(doc, tbl_lines)


def _build_week3_highlights_section(doc: Document) -> None:
    """Add highlights from Chapters 1 to 3 of the Viva Defense Guide."""
    _add_heading(doc, "2. Điểm Nhấn Chiến Lược & Lời Thoại (Chương 1–3)", level=2)
    _add_heading(doc, "Chương 1: Chiến thuật Báo cáo 'Nhả hàng' Thông minh", level=3)
    p1 = doc.add_paragraph(style="List Bullet")
    _add_inline_formatted_text(
        p1,
        "Khẳng định hoàn thành xuất sắc **120% khối lượng Giai đoạn 1**, "
        "tuyệt đối không tuyên bố hoàn tất 100% toàn bộ đề tài để tránh Scope Creep.",
    )
    p2 = doc.add_paragraph(style="List Bullet")
    _add_inline_formatted_text(
        p2,
        "Tạm ém bộ Benchmark 100 câu đối chứng khoa học để làm điểm nhấn bùng nổ cho Giữa kỳ và Cuối kỳ.",
    )

    _add_heading(doc, "Chương 2: Lời thoại Thuyết trình Tiến độ 5–7 Phút", level=3)
    items = [
        "**Phút 0–1:** Đặt vấn đề Temporal Misgrounding do Luật Thuế GTGT 48/2024 thay thế Luật 2008.",
        "**Phút 1–2.5:** Quy trình cào dữ liệu sạch HTML từ `vbpl.vn` và Parser 2 cấp độ.",
        "**Phút 2.5–4.5:** Mô hình Evolving Dual-Node trên Neo4j với 3 loại quan hệ AMENDS, GUIDES, CITES.",
        "**Phút 4.5–6:** Tìm kiếm lai SOTA kết hợp BM25 + Vector 1024 chiều và RRF (k=60).",
        "**Phút 6–7:** Demo nhanh Prototype Streamlit 30 giây và cam kết lộ trình Tuần 4–8.",
    ]
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        _add_inline_formatted_text(p, it)

    _add_heading(doc, "Chương 3: Hướng dẫn Trình diễn Trực tiếp (Live Demo)", level=3)
    p3 = doc.add_paragraph(style="List Bullet")
    _add_inline_formatted_text(
        p3,
        "Chuẩn bị sẵn 2 tab song song: **Neo4j Browser** (`localhost:7474`) và **Streamlit** (`localhost:8501`).",
    )
    p4 = doc.add_paragraph(style="List Bullet")
    _add_inline_formatted_text(
        p4,
        "Thao tác 2 ca kiểm chứng: Ca an toàn (15/08/2025 áp dụng Luật 48) và Ca bẫy thời gian (01/05/2025 lọc bỏ Luật 48).",
    )


def _build_week3_partner_brief_section(doc: Document) -> None:
    """Add slide presentation structure and collaboration plan."""
    _add_heading(doc, "3. Tóm Tắt Slide Cho Thành Viên Đồng Hành", level=2)
    slides = [
        "**Slide 1:** Bối cảnh thay đổi pháp lý 2024–2026 & Vấn nạn Temporal Misgrounding.",
        "**Slide 2:** Quy trình thu thập dữ liệu sạch & Thể thức lập pháp Việt Nam.",
        "**Slide 3:** Kiến trúc Đồ thị Tri thức Đa quan hệ 2 tầng & Evolving Dual-Node.",
        "**Slide 4:** Thuật toán Tìm kiếm lai SOTA (Hybrid Search & RRF k=60).",
        "**Slide 5:** Kết quả thực tế trên Prototype Streamlit Split-View.",
        "**Slide 6:** Lộ trình Sprint 4–8 (API FastAPI, React Frontend, Benchmark 100 câu).",
    ]
    for s in slides:
        p = doc.add_paragraph(style="List Bullet")
        _add_inline_formatted_text(p, s)


def _build_week3_cheat_sheet_section(doc: Document) -> None:
    """Add quick reference sheet with gold numbers and fail-safe Cypher queries."""
    _add_heading(doc, "4. Bảng Tra Cứu Nhanh Bỏ Túi (Cheat Sheet Tuần 3)", level=2)
    _add_heading(doc, "Các con số kỹ thuật 'Vàng'", level=3)
    p = doc.add_paragraph()
    _add_inline_formatted_text(
        p,
        "**5** văn bản | **76** Điều | **271** Khoản | **325** Điểm | "
        "**8** AMENDS | **3** GUIDES | **17** CITES | **< 5ms** Cypher latency.",
    )
    _add_heading(doc, "Ba câu lệnh Cypher cứu cánh", level=3)
    cypher_queries = [
        "// 1. Dem tong so node trong do thi",
        "MATCH (n) RETURN count(n);",
        "",
        "// 2. Kiem tra cac canh quan he sua doi AMENDS",
        "MATCH (a)-[r:AMENDS]->(b) RETURN a.id, r.action, b.id;",
        "",
        "// 3. Truy van lat cat thoi gian tai 15/08/2025",
        "MATCH (a:Article)",
        "WHERE a.valid_from <= '2025-08-15' <= a.valid_to",
        "RETURN a.id, a.title LIMIT 5;",
    ]
    _add_code_block(doc, cypher_queries)


def export_week3_summary_docx(
    input_md: str | Path, output_docx: str | Path
) -> Path:
    """Generate a compact Week 3 progress report Word document."""
    out_path = Path(output_docx)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _setup_document_styles(doc)
    _build_week3_summary_header(doc)
    _build_week3_goals_section(doc)
    _build_week3_highlights_section(doc)
    _build_week3_partner_brief_section(doc)
    _build_week3_cheat_sheet_section(doc)

    doc.save(str(out_path))
    return out_path


def main() -> None:
    """Execute exports for viva defense guide and week 3 summary."""
    convert_viva_guide_to_docx(
        input_md="docs/VIVA-DEFENSE-GUIDE.md",
        output_docx="docs/VIVA-DEFENSE-GUIDE.docx",
    )
    export_week3_summary_docx(
        input_md="docs/VIVA-DEFENSE-GUIDE.md",
        output_docx="docs/BAO-CAO-TIEN-DO-TUAN-3.docx",
    )


if __name__ == "__main__":
    main()
