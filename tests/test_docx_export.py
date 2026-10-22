from pathlib import Path
import docx
from scripts.export_viva_to_docx import (
    convert_viva_guide_to_docx,
    export_week3_summary_docx,
)


def test_convert_viva_guide_to_docx():
    input_md = "docs/VIVA-DEFENSE-GUIDE.md"
    output_docx = "docs/VIVA-DEFENSE-GUIDE.docx"
    result = convert_viva_guide_to_docx(input_md=input_md, output_docx=output_docx)

    out_file = Path(output_docx)
    assert out_file.exists(), f"Output file does not exist: {output_docx}"
    assert out_file.stat().st_size > 25 * 1024, "File size is under 25 KB"

    doc = docx.Document(output_docx)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for chapter in ["CHƯƠNG 1", "CHƯƠNG 2", "CHƯƠNG 3", "CHƯƠNG 4", "CHƯƠNG 5"]:
        assert any(chapter in h for h in headings), f"Missing heading: {chapter}"

    assert len(doc.paragraphs) > 100, f"Expected >100 paragraphs, got {len(doc.paragraphs)}"
    assert len(doc.tables) >= 3, f"Expected at least 3 tables, got {len(doc.tables)}"


def test_export_summary_report():
    input_md = "docs/VIVA-DEFENSE-GUIDE.md"
    output_docx = "docs/BAO-CAO-TIEN-DO-TUAN-3.docx"
    result = export_week3_summary_docx(input_md=input_md, output_docx=output_docx)

    out_file = Path(output_docx)
    assert out_file.exists(), f"Output file does not exist: {output_docx}"

    doc = docx.Document(output_docx)
    assert len(doc.paragraphs) > 0, "No paragraphs found in summary report"
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) > 0, "No headings found in summary report"
